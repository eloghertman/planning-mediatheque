# -*- coding: utf-8 -*-
"""
regeneration_sources.py
========================

Permet, lors d'une régénération partielle (Bloc 4), de déposer des fichiers
sources bruts MIS À JOUR (les mêmes types de fichiers que le Bloc 1 : congés,
accueil crèche, accueil de classe, lectures AssMat/AssPar, accueil CML
mercredis/vacances) pour corriger un événement déjà présent dans le planning
mais incomplet (ex: "Accueil classe" généré sans bibliothécaire à l'étape 1,
faute d'info à l'époque).

Un seul dépôt de fichiers, formats mélangés (.xlsx et .docx) : le TYPE de
chaque fichier est deviné automatiquement à partir de son contenu (même
principe que la détection par contenu déjà utilisée ailleurs dans l'outil —
plus robuste qu'un nom de fichier, qui varie d'un dépôt à l'autre).

Portée volontairement limitée (validée avec Elo, 09/2026) :
  - Un événement déjà affiché dans le planning, mais incomplet (jaune, sans
    agent) → corrigé PROPREMENT : le prénom est ajouté dans la case
    Accueil/Animation (ou Réunion/Absence), le jaune et le commentaire
    "incomplet" disparaissent, et une éventuelle note tapée à la main pour
    pallier ce manque (créant un doublon) est supprimée automatiquement si
    elle correspond clairement (même agent, même catégorie, horaire qui se
    recoupe).
  - Un événement TOUT À FAIT NOUVEAU (absent du planning avant même la
    régénération) : pris en compte pour que le calcul sache que l'agent
    concerné est occupé (il n'est donc pas proposé au comptoir à ce
    moment-là), mais PAS ajouté visuellement dans la colonne Accueil/
    Animation — à ajouter à la main via la zone de notes si besoin.
"""

import re

import openpyxl
from docx import Document as DocxDocument
from openpyxl.styles import PatternFill
from openpyxl.worksheet.formula import ArrayFormula

from sources_to_evenements import (
    MOIS_FR_CAP, normalize_agent,
    parse_conges, parse_accueil_creche, parse_accueil_classe,
    parse_lecture_jeudi_matin, parse_accueil_cml_mercredis,
    parse_accueil_cml_vacances,
)
from generate_planning_excel_septembre import (
    classer_evenement, label_evenement, fmt_hhmm,
    _extraire_litteral_reel, BAKED_LITERAL_RE, _notes_esc,
)
from planning_engine_cpsat import _parse_fr_time


TYPES_SOURCE_LABELS = {
    'conges': 'Congés équipe',
    'accueil_creche': 'Accueil libre crèche',
    'accueil_classe': 'Accueil de classe',
    'lecture_jeudi': 'Lectures AssMat/AssPar',
    'accueil_cml_mercredis': 'Accueil CML (mercredi)',
    'accueil_cml_vacances': 'Accueil CML (vacances)',
}

_PARSEURS = {
    'conges': parse_conges,
    'accueil_creche': parse_accueil_creche,
    'accueil_classe': parse_accueil_classe,
    'lecture_jeudi': parse_lecture_jeudi_matin,
    'accueil_cml_mercredis': parse_accueil_cml_mercredis,
    'accueil_cml_vacances': parse_accueil_cml_vacances,
}


# ══════════════════════════════════════════════════════════════
#  1. DÉTECTION DU TYPE DE FICHIER (contenu, pas nom de fichier)
# ══════════════════════════════════════════════════════════════

def detecter_type_source(path):
    """Devine le type d'un fichier source déposé, à partir de son extension
    et de sa structure interne. Renvoie une clé de TYPES_SOURCE_LABELS, ou
    None si non reconnu."""
    bas = path.lower()

    if bas.endswith('.docx'):
        try:
            doc = DocxDocument(path)
        except Exception:
            return None
        if not doc.tables:
            return None
        premiere_entete = [c.text.strip().lower() for c in doc.tables[0].rows[0].cells]
        if premiere_entete[:2] == ['date', 'cml']:
            return 'accueil_cml_mercredis'
        for table in doc.tables:
            entete = [c.text.strip().lower() for c in table.rows[0].cells]
            if 'centre' in entete and 'horaire' in entete:
                return 'accueil_cml_vacances'
        return None

    if bas.endswith(('.xlsx', '.xlsm')):
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception:
            return None
        ws = wb[wb.sheetnames[0]]

        # Congés équipe : un onglet nommé d'après un mois, avec "Nom de
        # l'employé" quelque part dans les premières lignes de cet onglet.
        for nom_onglet in wb.sheetnames:
            if nom_onglet.strip().capitalize() in MOIS_FR_CAP.values():
                ws_mois = wb[nom_onglet]
                for r in range(1, min(ws_mois.max_row, 15) + 1):
                    v = ws_mois.cell(row=r, column=2).value
                    if v and "nom de l" in str(v).lower():
                        return 'conges'

        a1 = ws['A1'].value
        heure_a1 = bool(a1) and re.match(r'\s*\d{1,2}\s*h', str(a1).lower())
        if heure_a1:
            # Accueil crèche (date en texte 'jeudi 7' colonne B) vs lecture
            # du jeudi matin (vraie date Excel colonne B) : on regarde le
            # type de la première cellule non vide de la colonne B.
            for r in range(2, min(ws.max_row, 30) + 1):
                b = ws.cell(row=r, column=2).value
                if b:
                    if hasattr(b, 'year'):  # datetime réel
                        return 'lecture_jeudi'
                    return 'accueil_creche'

        # Accueil de classe : "Nom de l'école" quelque part dans les 10
        # premières lignes.
        for r in range(1, min(ws.max_row, 10) + 1):
            for c in range(1, min(ws.max_column, 12) + 1):
                v = ws.cell(row=r, column=c).value
                if v and "nom de l'école" in str(v).lower().replace('’', "'"):
                    return 'accueil_classe'

        return None

    return None


# ══════════════════════════════════════════════════════════════
#  2. LECTURE, RESTREINTE AUX JOURS RÉGÉNÉRÉS
# ══════════════════════════════════════════════════════════════

def parser_fichiers_sources(chemins_et_types, jours_cibles_dates):
    """chemins_et_types : liste de (chemin, type_source).
    jours_cibles_dates : liste de `datetime.date` des jours à régénérer.

    Renvoie (evenements, stats) : evenements = liste d'événements (format
    sources_to_evenements) restreinte aux dates demandées ; stats = un
    résumé par fichier, pour affichage à l'utilisatrice."""
    mois_annees = sorted({(d.month, d.year) for d in jours_cibles_dates})
    dates_cibles = set(jours_cibles_dates)

    evenements = []
    stats = []
    for chemin, type_source in chemins_et_types:
        parseur = _PARSEURS.get(type_source)
        if parseur is None:
            stats.append({'chemin': chemin, 'type': type_source, 'trouves': 0,
                          'erreur': "Type de fichier non reconnu."})
            continue
        trouves = []
        erreur = None
        for mois, annee in mois_annees:
            try:
                trouves += parseur(chemin, mois, annee)
            except Exception as e:
                erreur = str(e)
                break
        if erreur:
            stats.append({'chemin': chemin, 'type': type_source, 'trouves': 0,
                          'erreur': erreur})
            continue
        pertinents = [ev for ev in trouves if ev['date'] in dates_cibles]
        evenements += pertinents
        stats.append({'chemin': chemin, 'type': type_source,
                      'trouves': len(pertinents), 'erreur': None})
    return evenements, stats


# ══════════════════════════════════════════════════════════════
#  3. CONTRAINTES POUR LE MOTEUR DE CALCUL (brique 2)
# ══════════════════════════════════════════════════════════════

def construire_contraintes_calcul(evenements_sources, jours_data_bruts):
    """Convertit les événements lus dans les fichiers sources mis à jour au
    même format que 'evenements_regeneres' / 'evenements_tous_jours' de
    regeneration_lecture.py (une entrée par (événement, agent)) — pour que
    le moteur sache que l'agent est occupé et ne le propose pas ailleurs."""
    contraintes = []
    for ev in evenements_sources:
        if not ev.get('agents'):
            continue  # toujours pas d'agent renseigné : rien à ajouter au calcul
        cs = _parse_fr_time(ev.get('debut'))
        ce = _parse_fr_time(ev.get('fin'))
        if cs is None or ce is None:
            continue
        date_str = ev['date'].isoformat()
        jour_data = next((jd for jd in jours_data_bruts.values()
                          if jd.get('date_str') == date_str), None)
        if jour_data is None:
            continue  # date hors de la semaine régénérée : ignorée
        type_evt = classer_evenement(ev['nom']).replace(' ', '')
        for agent in ev['agents']:
            agent_norm = normalize_agent(agent) or agent
            contraintes.append({
                'date': date_str, 'cs': cs, 'ce': ce,
                'nom': ev['nom'], 'nom_affichage': ev['nom'],
                'type': type_evt, 'agents': [agent_norm],
            })
    return contraintes


# ══════════════════════════════════════════════════════════════
#  4. CORRECTION DE L'AFFICHAGE (colonnes H/I/J + notes W-Z)
# ══════════════════════════════════════════════════════════════

def _segments(texte):
    return [s.strip() for s in (texte or '').split(';') if s.strip()]


def _nom_sans_agents(segment):
    """'Accueil classe (Stéphanie)' -> 'Accueil classe'."""
    return re.sub(r'\s*\([^)]*\)\s*$', '', segment).strip()


def _segment_sans_agent(segment):
    """Vrai si ce segment ne mentionne aucun agent connu entre parenthèses —
    que ce soit parce qu'il n'y a pas de parenthèses du tout, ou parce
    qu'elles ne contiennent qu'une précision d'horaire (ex. 'Accueil classe
    (10h15-10h45)', sans prénom)."""
    m = re.search(r'\(([^)]*)\)\s*$', segment)
    if not m:
        return True
    for morceau in m.group(1).split(','):
        if normalize_agent(morceau.strip()):
            return False
    return True


def _parse_note_texte(raw):
    """Reproduit en Python la même interprétation que les formules Excel des
    notes agents (W-Z) : catégorie, début, fin (en minutes). None si vide ou
    non reconnu."""
    if not raw or not str(raw).strip():
        return None
    s = str(raw).strip()
    low = s.lower()
    if any(k in low for k in ('réunion', 'reunion', 'rdv')):
        cat = 'Réunion'
    elif any(k in low for k in ('congé', 'conge', 'absen', 'part', 'formation')):
        cat = 'Absence'
    else:
        cat = 'Accueil/Animation'
    espace = s.find(' ')
    ft = s if espace == -1 else s[:espace]
    deb = fin = None
    if 'h' in ft.lower():
        if '-' in ft:
            a, b = ft.split('-', 1)
            deb, fin = _parse_fr_time(a), _parse_fr_time(b)
        else:
            deb = _parse_fr_time(ft)
            fin = deb + 1 if deb is not None else None
    return {'cat': cat, 'deb': deb, 'fin': fin}


COL_PAR_TYPE = {'Accueil/Animation': 8, 'Réunion': 9, 'Absence': 10}


def appliquer_corrections_affichage(ws, jour_data, evenements_sources):
    """Pour chaque événement lu dans les fichiers sources mis à jour qui
    correspond à un événement déjà affiché dans le planning MAIS encore
    incomplet (jaune, sans agent) : ajoute le prénom dans la case, retire le
    jaune + le commentaire "incomplet", et supprime une éventuelle note
    manuelle devenue redondante. Renvoie une liste de messages (str) pour
    affichage à l'utilisatrice."""
    messages = []
    for ev in evenements_sources:
        if not ev.get('agents'):
            continue
        cs = _parse_fr_time(ev.get('debut'))
        ce = _parse_fr_time(ev.get('fin'))
        if cs is None or ce is None:
            continue
        type_evt = classer_evenement(ev['nom']).replace(' ', '')
        col = COL_PAR_TYPE[type_evt]

        cren_cible = next((c for c in jour_data['creneaux']
                            if c['debut'] <= cs < c['fin']), None)
        if cren_cible is None:
            continue
        row = cren_cible['row']
        cell = ws.cell(row=row, column=col)

        litteral, ok = _extraire_litteral_reel(cell.value)
        if not ok or not litteral:
            continue

        nom_norm = ev['nom'].strip().lower()
        segments = _segments(litteral)
        idx_cible = next(
            (i for i, seg in enumerate(segments)
             if _nom_sans_agents(seg).lower() == nom_norm and _segment_sans_agent(seg)),
            None,
        )
        if idx_cible is None:
            continue  # rien à corriger dans cette case (déjà attribué, ou absent)

        ancien_segment = segments[idx_cible]
        nouveau_segment = label_evenement(
            {'nom': ev['nom'], 'agents': ev['agents'], 'cs': cs, 'ce': ce},
            cren_cible['debut'], cren_cible['fin'],
        )
        segments[idx_cible] = nouveau_segment
        nouveau_litteral = '; '.join(segments)

        texte_formule = cell.value.text if isinstance(cell.value, ArrayFormula) else cell.value
        m = BAKED_LITERAL_RE.match(texte_formule or '')
        if not m:
            continue
        nouveau_prefixe = f'=TRIM("{_notes_esc(nouveau_litteral)}"&IF('
        cell.value = ArrayFormula(ref=cell.coordinate, text=nouveau_prefixe + texte_formule[m.end():])

        if len(segments) == 1:
            # Seul événement de la case : jaune + commentaire "incomplet" retirés.
            cell.fill = PatternFill(fill_type=None)
            if cell.comment and str(cell.comment.text).startswith('INCOMPLET'):
                cell.comment = None

        heure_txt = f"{fmt_hhmm(cren_cible['debut'])}-{fmt_hhmm(cren_cible['fin'])}"
        messages.append(f"« {ancien_segment} » complété en « {nouveau_segment} » ({heure_txt}).")

        # ── Suppression d'une note manuelle devenue redondante ──────────
        agents_norm = {normalize_agent(a) or a for a in ev['agents']}
        for r in range(jour_data['row_debut_data'], jour_data['row_fin_data'] + 1):
            for col_nom, col_evt in ((23, 24), (25, 26)):
                nom_agent_cell = ws.cell(row=r, column=col_nom).value
                if not nom_agent_cell or str(nom_agent_cell).strip() not in agents_norm:
                    continue
                note_cell = ws.cell(row=r, column=col_evt)
                parsed = _parse_note_texte(note_cell.value)
                if not parsed or parsed['cat'] != type_evt:
                    continue
                if parsed['deb'] is None or parsed['fin'] is None:
                    continue
                if cs < parsed['fin'] and ce > parsed['deb']:  # chevauchement
                    messages.append(
                        f"Note manuelle de {nom_agent_cell} devenue inutile "
                        f"supprimée : « {note_cell.value} »."
                    )
                    note_cell.value = None
    return messages
